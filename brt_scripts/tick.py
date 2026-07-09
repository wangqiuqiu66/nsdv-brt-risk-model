"""Run tick BRT models with Set 4 weight and RC >= 0.02 screening.

Includes RC-table and high-risk area/population table export.
"""

from __future__ import annotations

import os
from pathlib import Path

import numpy as np
import pandas as pd
import xgboost as xgb
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from sklearn.metrics import auc, roc_curve
from sklearn.model_selection import train_test_split

# ---- inline helpers (formerly brt_utils) ----

def _feature_importance_frame(model) -> pd.DataFrame:
    importance = model.get_booster().get_score(importance_type="weight")
    frame = pd.DataFrame(list(importance.items()), columns=["Key", "Value"])
    frame["rc"] = frame["Value"] / frame["Value"].sum()
    return frame.sort_values("rc", ascending=False).reset_index(drop=True)


def _choose_key_factors(importance_df: pd.DataFrame, min_rc: float) -> list[str]:
    return importance_df.loc[importance_df["rc"] >= min_rc, "Key"].tolist()


def _summarize_metrics(metrics_df: pd.DataFrame, label: str) -> pd.DataFrame:
    stats = {"species": label}
    for metric in ["auc", "tpr", "tnr", "threshold"]:
        if metric in metrics_df.columns and not metrics_df[metric].empty:
            values = metrics_df[metric]
            stats[f"{metric}_mean"] = values.mean()
            stats[f"{metric}_lbd"] = np.percentile(values, 2.5)
            stats[f"{metric}_ubd"] = np.percentile(values, 97.5)
    return pd.DataFrame([stats]).round(3)


def _summarize_predictions(pred_df: pd.DataFrame) -> pd.DataFrame:
    return (
        pred_df.groupby("gaul2_code")["pred"]
        .agg(
            mean_prob="mean",
            ci_lower=lambda x: np.percentile(x, 2.5),
            ci_upper=lambda x: np.percentile(x, 97.5),
        )
        .reset_index()
    )


PROJECT_ROOT = Path(__file__).resolve().parents[1]
INPUT_DIR = PROJECT_ROOT / "input"
OUTPUT_BASE_ROOT = Path(os.environ.get("BRT_OUTPUT_DIR", PROJECT_ROOT / "output"))
OUTPUT_ROOT = OUTPUT_BASE_ROOT / "host"
RC_TABLE_DIR = OUTPUT_BASE_ROOT / "tick_rc_tables"

RISK_DATA_PATH = INPUT_DIR / "risk_data_all.xlsx"
TICK_DIST_PATH = INPUT_DIR / "tick_location.xlsx"
WEIGHT_PATH = INPUT_DIR / "weight.xlsx"

TARGET_SPECIES_LIST = [
    "Rhipicephalus appendiculatus",
    "Haemaphysalis intermedia",
    "Haemaphysalis longicornis",
]

SCREEN_ITERATIONS = 10
N_ITERATIONS = 100
NEG_SAMPLE_RATIO = 5
KEY_FACTOR_MIN_RC = 0.02
HIGH_RISK_THRESHOLD = 0.5

NON_ENVIRONMENT_COLUMNS = {
    "gaul2_code", "gaul0_name", "gaul1_name", "gaul2_name", "continent",
    "status", "model_weight", "rescaled", "prob", "recip",
}

FIXED_PARAMS = {
    "n_estimators": 1000, "max_depth": 5, "learning_rate": 0.005,
    "subsample": 0.75, "objective": "binary:logistic", "eval_metric": "logloss",
    "n_jobs": 1,
}

SPECIES_HEADERS = {
    "Rhipicephalus appendiculatus": "R. appendiculatus",
    "Haemaphysalis intermedia": "H. intermedia",
    "Haemaphysalis longicornis": "H. longicornis",
}

RISK_FACTOR_LABELS = {
    "elevation": "Elevation",
    "Bio1": "Bio1", "Bio2": "Bio2", "Bio3": "Bio3", "Bio4": "Bio4",
    "Bio5": "Bio5", "Bio6": "Bio6", "Bio7": "Bio7", "Bio8": "Bio8",
    "Bio9": "Bio9", "Bio10": "Bio10", "Bio11": "Bio11", "Bio12": "Bio12",
    "Bio13": "Bio13", "Bio14": "Bio14", "Bio15": "Bio15", "Bio16": "Bio16",
    "Bio17": "Bio17", "Bio18": "Bio18", "Bio19": "Bio19",
    "LandCoverType_10": "Tree cover",
    "LandCoverType_20": "Shrubland", "LandCoverType_30": "Grassland",
    "LandCoverType_40": "Cropland", "LandCoverType_50": "Built-up",
    "LandCoverType_60": "Bare/sparse vegetation", "LandCoverType_70": "Snow and ice",
    "LandCoverType_80": "Permanent water bodies", "LandCoverType_90": "Herbaceous wetland",
    "LandCoverType_95": "Mangroves", "LandCoverType_100": "Moss and lichen",
    "Buffaloes_density": "Buffaloes", "Sheep_density": "Sheep",
    "Goats_density": "Goats", "Cattle_density": "Cattle",
}


# ======================================================================
#  model functions
# ======================================================================

def get_environment_features(df: pd.DataFrame) -> list[str]:
    return [
        col for col in df.columns
        if col not in NON_ENVIRONMENT_COLUMNS
        and pd.api.types.is_numeric_dtype(df[col])
    ]


def add_model_weights(full_risk_data: pd.DataFrame) -> pd.DataFrame:
    weight_df = pd.read_excel(WEIGHT_PATH)

    weighted = full_risk_data.copy()
    weighted["gaul2_code"] = weighted["gaul2_code"].astype(str)
    weight_df = weight_df[["gaul2_code", "rescaled"]].copy()
    weight_df["gaul2_code"] = weight_df["gaul2_code"].astype(str)

    weighted = weighted.merge(weight_df, on="gaul2_code", how="left")
    weighted["model_weight"] = weighted["rescaled"].fillna(1.0)
    return weighted.drop(columns=["rescaled"])


def screen_key_factors(
    pos_data: pd.DataFrame, background_pool: pd.DataFrame,
    n_neg_needed: int, feature_cols: list[str],
) -> tuple[list[str], pd.DataFrame]:
    importance_frames = []

    for iteration in range(1, SCREEN_ITERATIONS + 1):
        neg_sample = background_pool.sample(n=n_neg_needed, random_state=iteration, replace=False)
        data = pd.concat([pos_data, neg_sample], ignore_index=True)
        train, _ = train_test_split(data, test_size=0.25, random_state=iteration, stratify=data["status"])

        n_train_neg = (train["status"] == 0).sum()
        n_train_pos = (train["status"] == 1).sum()
        params = FIXED_PARAMS.copy()
        params["scale_pos_weight"] = n_train_neg / n_train_pos

        model = xgb.XGBClassifier(random_state=iteration, **params)
        model.fit(train[feature_cols], train["status"], sample_weight=train["model_weight"])

        importance = _feature_importance_frame(model)
        importance["iteration"] = iteration
        importance_frames.append(importance)

    rc_df = pd.concat(importance_frames, ignore_index=True)
    if rc_df.empty or rc_df["Value"].sum() == 0:
        rc_grouped = pd.DataFrame(columns=["Key", "sum_value", "rc"])
        return [], rc_grouped

    rc_grouped = rc_df.groupby("Key")["Value"].sum().reset_index(name="sum_value")
    rc_grouped["rc"] = rc_grouped["sum_value"] / rc_grouped["sum_value"].sum()
    key_factors = _choose_key_factors(rc_grouped, KEY_FACTOR_MIN_RC)
    return key_factors, rc_grouped.sort_values("rc", ascending=False).reset_index(drop=True)


def run_species_workflow(
    species_name: str, full_risk_data_raw: pd.DataFrame, tick_loc_df: pd.DataFrame,
) -> None:
    output_dir = OUTPUT_ROOT / species_name / "unified_result"
    output_metric_dir = output_dir / "metric"
    output_rc_dir = output_dir / "rc"
    output_pred_dir = output_dir / "prediction"
    output_hrc_dir = output_dir / "hrc"
    for p in [output_metric_dir, output_rc_dir, output_pred_dir, output_hrc_dir]:
        p.mkdir(parents=True, exist_ok=True)

    full_risk_data = add_model_weights(full_risk_data_raw)
    feature_cols = get_environment_features(full_risk_data)

    target_tick_rows = tick_loc_df[tick_loc_df["Species"] == species_name]
    pos_gaul_codes = set(target_tick_rows["gaul2_code"].dropna().astype(str).unique())
    full_risk_data["status"] = full_risk_data["gaul2_code"].isin(pos_gaul_codes).astype(int)

    pos_data = full_risk_data[full_risk_data["status"] == 1].copy()
    background_pool = full_risk_data[full_risk_data["status"] == 0].copy()
    n_neg_needed = min(len(background_pool), len(pos_data) * NEG_SAMPLE_RATIO)

    # step 1 鈥?screen key factors
    key_factors, screening_rc = screen_key_factors(pos_data, background_pool, n_neg_needed, feature_cols)
    screening_rc.to_csv(output_rc_dir / f"{species_name}_rc_screening.csv", index=False)

    # step 2 鈥?100-iteration bootstrap prediction
    metrics_list, all_rc_dfs, all_hrc_dfs, all_pred_dfs = [], [], [], []

    for iteration in range(1, N_ITERATIONS + 1):
        neg_sample = background_pool.sample(n=n_neg_needed, random_state=iteration, replace=False)
        data = pd.concat([pos_data, neg_sample], ignore_index=True)
        train, test = train_test_split(data, test_size=0.25, random_state=iteration, stratify=data["status"])

        n_train_neg = (train["status"] == 0).sum()
        n_train_pos = (train["status"] == 1).sum()
        current_params = FIXED_PARAMS.copy()
        current_params["scale_pos_weight"] = n_train_neg / n_train_pos

        model = xgb.XGBClassifier(random_state=iteration, **current_params)
        model.fit(train[key_factors], train["status"], sample_weight=train["model_weight"])

        # metrics
        prob_test = model.predict_proba(test[key_factors])[:, 1]
        fpr, tpr, thresholds = roc_curve(test["status"], prob_test)
        optimal_idx = int(np.argmax(tpr - fpr))
        metrics_list.append({
            "iteration": iteration, "auc": auc(fpr, tpr),
            "tpr": tpr[optimal_idx], "tnr": 1 - fpr[optimal_idx],
            "threshold": thresholds[optimal_idx], "n_features_final": len(key_factors),
        })

        # RC
        rc_final = _feature_importance_frame(model)
        rc_final["iteration"] = iteration
        all_rc_dfs.append(rc_final)

        # predictions
        pred_probs = model.predict_proba(full_risk_data[key_factors])[:, 1]
        pred_df = pd.DataFrame({
            "gaul2_code": full_risk_data["gaul2_code"],
            "round": iteration, "status": full_risk_data["status"], "pred": pred_probs,
        })
        all_pred_dfs.append(pred_df)

        # high-risk counties
        hrc_df = pred_df[pred_df["pred"] >= HIGH_RISK_THRESHOLD].copy()
        hrc_df["threshold"] = HIGH_RISK_THRESHOLD
        hrc_df["iteration"] = iteration
        if not hrc_df.empty:
            all_hrc_dfs.append(hrc_df)
            hrc_df.to_csv(output_hrc_dir / f"{species_name}_hrc_iter_{iteration}.csv", index=False)

    # write output files
    metrics_df = pd.DataFrame(metrics_list)
    metrics_df.to_csv(output_metric_dir / f"{species_name}_metrics_raw.csv", index=False)
    _summarize_metrics(metrics_df, species_name).to_csv(
        output_metric_dir / f"{species_name}_metrics_statistics.csv", index=False)

    full_rc_df = pd.concat(all_rc_dfs, ignore_index=True)
    rc_grouped = full_rc_df.groupby("Key")["Value"].sum().reset_index(name="sum_value")
    rc_grouped = (
        pd.DataFrame({"Key": feature_cols})
        .merge(rc_grouped, on="Key", how="left")
        .fillna({"sum_value": 0})
    )
    rc_grouped["rc"] = (rc_grouped["sum_value"] / rc_grouped["sum_value"].sum()).round(4)
    rc_grouped.sort_values("rc", ascending=False).to_csv(
        output_rc_dir / f"{species_name}_rc_aggregated.csv", index=False)

    if all_hrc_dfs:
        pd.concat(all_hrc_dfs, ignore_index=True).to_csv(
            output_hrc_dir / f"{species_name}_hrc_all_iterations.csv", index=False)

    pred_all = pd.concat(all_pred_dfs, ignore_index=True)
    _summarize_predictions(pred_all).to_csv(
        output_pred_dir / f"{species_name}_prediction_summary.csv",
        index=False, float_format="%.4f")
    pred_all.pivot(index="gaul2_code", columns="round", values="pred").reset_index().to_csv(
        output_pred_dir / f"{species_name}_full_prediction_matrix.csv",
        index=False, float_format="%.4f")

    pd.DataFrame({"key_factors": key_factors}).to_csv(output_dir / "key_factors.csv", index=False)


# ======================================================================
#  RC table export
# ======================================================================

def export_tick_rc_tables() -> None:
    mean_rc_frames = []
    for species_name in TARGET_SPECIES_LIST:
        rc_path = (
            OUTPUT_ROOT / species_name / "unified_result" / "rc"
            / f"{species_name}_rc_aggregated.csv"
        )
        df = pd.read_csv(rc_path)
        df = df.rename(columns={"Key": "risk_factor", "rc": "mean_rc"})
        df["species"] = species_name
        mean_rc_frames.append(df[["species", "risk_factor", "mean_rc"]])

    mean_rc = pd.concat(mean_rc_frames, ignore_index=True)
    table = mean_rc.pivot(index="risk_factor", columns="species", values="mean_rc")

    ordered = [f for f in RISK_FACTOR_LABELS if f in table.index] + [
        f for f in table.index if f not in RISK_FACTOR_LABELS]
    table = table.reindex(ordered)
    table = table.rename(index=RISK_FACTOR_LABELS, columns=SPECIES_HEADERS)
    table = table[[SPECIES_HEADERS[s] for s in TARGET_SPECIES_LIST]]
    table.insert(0, "Risk factor", table.index)
    table = table.reset_index(drop=True)
    for col in SPECIES_HEADERS.values():
        table[col] = table[col].map(lambda v: "" if pd.isna(v) else f"{float(v):.4f}")

    RC_TABLE_DIR.mkdir(parents=True, exist_ok=True)
    table.to_csv(RC_TABLE_DIR / "tick_relative_contribution_table.csv", index=False, encoding="utf-8-sig")

    xlsx_path = RC_TABLE_DIR / "tick_relative_contribution_tables.xlsx"
    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        table.to_excel(writer, sheet_name="RC table", index=False)

    wb = load_workbook(xlsx_path)
    ws = wb["RC table"]
    ws.insert_rows(1)
    ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=4)
    ws.cell(1, 1, "Risk factor")
    ws.cell(1, 2, "Tick species")
    thin = Side(style="thin", color="000000")
    bdr = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=4):
        for cell in row:
            cell.border = bdr
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.font = Font(name="Times New Roman", size=12)
    for cell in ws[1] + ws[2]:
        cell.font = Font(name="Times New Roman", size=12, bold=True, italic=cell.column > 1)
    for col_idx in range(1, 5):
        ws.column_dimensions[get_column_letter(col_idx)].width = 26 if col_idx == 1 else 22
    for row_idx in range(1, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = 24
    wb.save(xlsx_path)

    print(f"Saved RC table: {RC_TABLE_DIR / 'tick_relative_contribution_table.csv'}")
    print(f"Saved RC workbook: {xlsx_path}")


# ======================================================================
#  high-risk area / population summary tables
# ======================================================================
POPULATION_PATH = INPUT_DIR / "population.xlsx"
HIGH_RISK_DIR_TICK = OUTPUT_BASE_ROOT / "high_risk_tables"


def _load_county_attrs() -> pd.DataFrame:
    names = pd.read_excel(RISK_DATA_PATH, usecols=["gaul2_code", "gaul0_name", "gaul1_name", "gaul2_name"])
    names["gaul2_code"] = names["gaul2_code"].astype(str)
    names = names.drop_duplicates("gaul2_code")
    pop = pd.read_excel(POPULATION_PATH)
    pop["gaul2_code"] = pop["gaul2_code"].astype(str)
    pop = pop[["gaul2_code", "area (km²)", "population density (/km²)"]].copy()
    pop = pop.rename(columns={"area (km²)": "area_km2", "population density (/km²)": "population_density"})
    pop["population"] = (pop["area_km2"] * pop["population_density"]).fillna(0).round().astype("int64")
    attrs = names.merge(pop, on="gaul2_code", how="left")
    for c in ["area_km2", "population_density"]:
        attrs[c] = attrs[c].fillna(0)
    attrs["population"] = attrs["population"].fillna(0).astype("int64")
    return attrs


def _summarize_codes(attrs: pd.DataFrame, codes: pd.Series, label: str, range_type: str) -> dict:
    codes = codes.dropna().astype(str).drop_duplicates()
    subset = attrs[attrs["gaul2_code"].isin(codes)]
    area = float(subset["area_km2"].sum())
    pop = int(subset["population"].sum())
    return {
        "label": label, "range_type": range_type,
        "county_count": int(subset.shape[0]),
        "country_count": int(subset["gaul0_name"].nunique()),
        "area_km2": area, "area_million_km2": area / 1e6,
        "population": pop, "population_million": pop / 1e6,
    }


def export_tick_high_risk_tables() -> None:
    attrs = _load_county_attrs()
    tick_loc = pd.read_excel(TICK_DIST_PATH)
    tick_loc["gaul2_code"] = tick_loc["gaul2_code"].astype(str)
    summary_rows = []

    for species_name in TARGET_SPECIES_LIST:
        pred_path = (
            OUTPUT_ROOT / species_name / "unified_result" / "prediction"
            / f"{species_name}_prediction_summary.csv"
        )
        pred = pd.read_csv(pred_path)
        pred["gaul2_code"] = pred["gaul2_code"].astype(str)
        prob_col = "mean_prob" if "mean_prob" in pred.columns else "mean"
        high_risk = pred[pred[prob_col] >= 0.5]

        obs_codes = tick_loc.loc[tick_loc["Species"].eq(species_name), "gaul2_code"]
        summary_rows.append(_summarize_codes(attrs, obs_codes, species_name, "observed"))
        summary_rows.append(_summarize_codes(attrs, high_risk["gaul2_code"], species_name, "predicted_high_risk"))

    summary = pd.DataFrame(summary_rows)

    comp_rows = []
    metrics = ["county_count", "country_count", "area_km2", "area_million_km2", "population", "population_million"]
    for (species_name,), group in summary.groupby(["label"], sort=False):
        obs = group[group["range_type"].eq("observed")].iloc[0]
        pred = group[group["range_type"].eq("predicted_high_risk")].iloc[0]
        row = {"label": species_name}
        for m in metrics:
            row[f"observed_{m}"] = obs[m]
            row[f"predicted_{m}"] = pred[m]
            row[f"increase_{m}"] = pred[m] - obs[m]
            row[f"increase_pct_{m}"] = 100 * (pred[m] - obs[m]) / obs[m] if obs[m] else pd.NA
        comp_rows.append(row)
    comparison = pd.DataFrame(comp_rows)

    HIGH_RISK_DIR_TICK.mkdir(parents=True, exist_ok=True)
    summary.to_csv(HIGH_RISK_DIR_TICK / "tick_high_risk_summary.csv", index=False, float_format="%.6f")
    comparison.to_csv(HIGH_RISK_DIR_TICK / "tick_high_risk_increase_vs_observed.csv", index=False, float_format="%.6f")
    with pd.ExcelWriter(HIGH_RISK_DIR_TICK / "tick_high_risk_areas.xlsx") as w:
        summary.to_excel(w, sheet_name="summary", index=False)
        comparison.to_excel(w, sheet_name="increase_vs_observed", index=False)

    print(f"Saved tick high-risk tables: {HIGH_RISK_DIR_TICK}")
    print(summary.to_string(index=False))


# ======================================================================
#  main
# ======================================================================

def main() -> None:
    full_risk_data = pd.read_excel(RISK_DATA_PATH)
    tick_loc_df = pd.read_excel(TICK_DIST_PATH)

    for species_name in TARGET_SPECIES_LIST:
        run_species_workflow(species_name, full_risk_data, tick_loc_df)

    # RC tables
    export_tick_rc_tables()

    # high-risk area / population summary tables
    export_tick_high_risk_tables()


if __name__ == "__main__":
    main()


